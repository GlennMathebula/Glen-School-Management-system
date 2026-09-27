from io import BytesIO

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


# ============================================================
# EXTRACTION LIMIT
# ============================================================

MAX_EXTRACTED_TEXT_CHARACTERS = (
    250_000
)


# ============================================================
# CLEAN EXTRACTED TEXT
# ============================================================

def clean_extracted_text(
    text_value: str,
) -> str:

    if not text_value:
        return ""

    lines = []

    for line in (
        text_value
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .split("\n")
    ):

        cleaned_line = (
            " ".join(
                line
                .strip()
                .split()
            )
        )

        if cleaned_line:
            lines.append(
                cleaned_line
            )

    text_value = (
        "\n".join(
            lines
        )
        .strip()
    )

    if (
        len(
            text_value
        )
        > MAX_EXTRACTED_TEXT_CHARACTERS
    ):

        text_value = (
            text_value[
                :MAX_EXTRACTED_TEXT_CHARACTERS
            ]
        )

    return text_value


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(
    file_bytes: bytes,
) -> str:

    try:

        reader = PdfReader(
            BytesIO(
                file_bytes
            )
        )

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):

            page_text = (
                page.extract_text()
                or ""
            )

            page_text = (
                clean_extracted_text(
                    page_text
                )
            )

            if page_text:

                pages.append(
                    (
                        f"PAGE {page_number}\n"
                        f"{page_text}"
                    )
                )

        return (
            clean_extracted_text(
                "\n\n".join(
                    pages
                )
            )
        )

    except Exception as error:

        raise ValueError(
            "The PDF could not be read."
        ) from error


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(
    file_bytes: bytes,
) -> str:

    try:

        document = Document(
            BytesIO(
                file_bytes
            )
        )

        content = []

        # ----------------------------------------------------
        # PARAGRAPHS
        # ----------------------------------------------------

        for paragraph in (
            document.paragraphs
        ):

            paragraph_text = (
                clean_extracted_text(
                    paragraph.text
                )
            )

            if paragraph_text:

                content.append(
                    paragraph_text
                )

        # ----------------------------------------------------
        # TABLES
        # ----------------------------------------------------

        for table_number, table in enumerate(
            document.tables,
            start=1,
        ):

            table_lines = []

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    cell_text = (
                        clean_extracted_text(
                            cell.text
                        )
                    )

                    cells.append(
                        cell_text
                    )

                if any(
                    cells
                ):

                    table_lines.append(
                        " | ".join(
                            cells
                        )
                    )

            if table_lines:

                content.append(
                    (
                        f"TABLE {table_number}\n"
                        + "\n".join(
                            table_lines
                        )
                    )
                )

        return (
            clean_extracted_text(
                "\n\n".join(
                    content
                )
            )
        )

    except Exception as error:

        raise ValueError(
            "The DOCX file could not be read."
        ) from error


# ============================================================
# PPTX TEXT EXTRACTION
# ============================================================

def extract_pptx_text(
    file_bytes: bytes,
) -> str:

    try:

        presentation = Presentation(
            BytesIO(
                file_bytes
            )
        )

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            slide_content = []

            for shape in slide.shapes:

                if hasattr(
                    shape,
                    "text",
                ):

                    shape_text = (
                        clean_extracted_text(
                            shape.text
                        )
                    )

                    if shape_text:

                        slide_content.append(
                            shape_text
                        )

                if getattr(
                    shape,
                    "has_table",
                    False,
                ):

                    table = (
                        shape.table
                    )

                    for row in (
                        table.rows
                    ):

                        row_text = []

                        for cell in (
                            row.cells
                        ):

                            cell_text = (
                                clean_extracted_text(
                                    cell.text
                                )
                            )

                            row_text.append(
                                cell_text
                            )

                        if any(
                            row_text
                        ):

                            slide_content.append(
                                " | ".join(
                                    row_text
                                )
                            )

            if slide_content:

                slides.append(
                    (
                        f"SLIDE {slide_number}\n"
                        + "\n".join(
                            slide_content
                        )
                    )
                )

        return (
            clean_extracted_text(
                "\n\n".join(
                    slides
                )
            )
        )

    except Exception as error:

        raise ValueError(
            "The PPTX file could not be read."
        ) from error


# ============================================================
# EXTRACT RESOURCE TEXT
# ============================================================

def extract_learning_resource_text(
    *,
    filename: str,
    mime_type: str | None,
    file_bytes: bytes,
) -> str:

    filename_lower = (
        filename
        .strip()
        .lower()
    )

    mime_type = (
        mime_type
        or ""
    ).strip().lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if (
        filename_lower.endswith(
            ".pdf"
        )
        or mime_type
        == "application/pdf"
    ):

        extracted_text = (
            extract_pdf_text(
                file_bytes
            )
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif (
        filename_lower.endswith(
            ".docx"
        )
        or mime_type
        == (
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml."
            "document"
        )
    ):

        extracted_text = (
            extract_docx_text(
                file_bytes
            )
        )

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    elif (
        filename_lower.endswith(
            ".pptx"
        )
        or mime_type
        == (
            "application/vnd.openxmlformats-"
            "officedocument.presentationml."
            "presentation"
        )
    ):

        extracted_text = (
            extract_pptx_text(
                file_bytes
            )
        )

    else:

        raise ValueError(
            "AI note generation currently "
            "supports PDF, DOCX and PPTX "
            "resources only."
        )

    # --------------------------------------------------------
    # EMPTY / SCANNED DOCUMENT
    # --------------------------------------------------------

    if not extracted_text.strip():

        raise ValueError(
            "No readable text could be "
            "extracted from this resource. "
            "The document may be scanned "
            "or image-based."
        )

    return extracted_text